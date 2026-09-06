import streamlit as st
import pandas as pd
from sqlalchemy import create_engine, text
from dotenv import load_dotenv
import os
import io
import hashlib
import hmac
import time
import base64
from datetime import datetime, timedelta
import extra_streamlit_components as stx

load_dotenv()

st.set_page_config(page_title="Worker Management", layout="wide")

# ============================================================
# AUTH
# ============================================================
_USERS = {
    "admin": hashlib.sha256("gofo1234".encode()).hexdigest(),
}
_SESSION_SECRET = os.getenv("SESSION_SECRET", "gofo_wm_session_key_2024")
_COOKIE_NAME    = "wm_session"
_SESSION_TTL    = 86400  # 24 hours in seconds

def _check_login(username: str, password: str) -> bool:
    return _USERS.get(username) == hashlib.sha256(password.encode()).hexdigest()

def _make_token(username: str) -> str:
    expiry  = str(int(time.time()) + _SESSION_TTL)
    payload = f"{username}:{expiry}"
    sig     = hmac.new(_SESSION_SECRET.encode(), payload.encode(), "sha256").hexdigest()
    return base64.urlsafe_b64encode(f"{payload}:{sig}".encode()).decode()

def _verify_token(token: str):
    try:
        decoded = base64.urlsafe_b64decode(token.encode()).decode()
        *parts, sig = decoded.split(":")
        username, expiry_str = parts[0], parts[1]
        payload      = f"{username}:{expiry_str}"
        expected_sig = hmac.new(_SESSION_SECRET.encode(), payload.encode(), "sha256").hexdigest()
        if not hmac.compare_digest(sig, expected_sig):
            return None
        if int(expiry_str) < int(time.time()):
            return None
        return username
    except Exception:
        return None

_cookie_mgr = stx.CookieManager()

# Auto-login from cookie
if not st.session_state.get("authenticated"):
    _token = _cookie_mgr.get(_COOKIE_NAME)
    if _token:
        _cookie_user = _verify_token(_token)
        if _cookie_user:
            st.session_state["authenticated"] = True
            st.session_state["username"]      = _cookie_user

if not st.session_state.get("authenticated"):
    st.title("Worker Management 排班管理")
    st.subheader("Sign in to continue")
    with st.form("login_form"):
        _user = st.text_input("Username")
        _pass = st.text_input("Password", type="password")
        if st.form_submit_button("Log In", use_container_width=True):
            if _check_login(_user, _pass):
                st.session_state["authenticated"] = True
                st.session_state["username"]      = _user
                _cookie_mgr.set(
                    _COOKIE_NAME,
                    _make_token(_user),
                    expires_at=datetime.now() + timedelta(seconds=_SESSION_TTL),
                )
                st.rerun()
            else:
                st.error("Invalid username or password.")
    st.stop()
PGHOST     = os.getenv("PGHOST")
PGPORT     = os.getenv("PGPORT")
PGUSER     = os.getenv("PGUSER")
PGPASSWORD = os.getenv("PGPASSWORD")
PGDATABASE = os.getenv("PGDATABASE")

engine = create_engine(
    f"postgresql://{PGUSER}:{PGPASSWORD}@{PGHOST}:{PGPORT}/{PGDATABASE}"
)

S = "worker_management"


def load_distinct(col):
    with engine.connect() as conn:
        rows = conn.execute(text(
            f"SELECT DISTINCT {col} FROM {S}.dim_labor WHERE {col} IS NOT NULL AND {col} != '' ORDER BY {col}"
        )).fetchall()
    return [""] + [r[0] for r in rows]

with st.sidebar:
    st.markdown(f"Signed in as **{st.session_state.get('username', '')}**")
    if st.button("Log Out", use_container_width=True):
        _cookie_mgr.delete(_COOKIE_NAME)
        st.session_state.clear()
        st.rerun()
    st.divider()
    st.markdown("**View**")
    st.radio(
        "Role",
        ["Sorting Lead", "Clerk", "Manager"],
        index=0,
        key="view_role",
        label_visibility="collapsed",
    )

st.title("Worker Management 排班管理")

st.markdown("""
<style>
div:has(span.del-marker) + div button {
    background-color: #e53935 !important;
    color: #ffffff !important;
    border: 1px solid #c62828 !important;
    font-size: 0.85em;
}
div:has(span.del-marker) + div button:hover {
    background-color: #b71c1c !important;
    border-color: #b71c1c !important;
}
</style>
""", unsafe_allow_html=True)

_view_role = st.session_state.get("view_role", "Sorting Lead")

_ROLE_TABS = {
    "Sorting Lead": ["Schedules 排班", "Area Analysis 区域分析", "Exceptions 异常管理", "Day Off 调休"],
    "Clerk":        ["Workers 劳务", "Attendance Export 签到表导出", "Temp Workers 临时工", "Exceptions 异常管理", "Day Off 调休"],
    "Manager":      ["Shifts 班次"],
}
_active_names = _ROLE_TABS[_view_role]
_tab_refs = dict(zip(_active_names, st.tabs(_active_names)))

tab_workers   = _tab_refs.get("Workers 劳务")
tab_shifts    = _tab_refs.get("Shifts 班次")
tab_schedules = _tab_refs.get("Schedules 排班")
tab_area      = _tab_refs.get("Area Analysis 区域分析")
tab_export    = _tab_refs.get("Attendance Export 签到表导出")
tab_temp      = _tab_refs.get("Temp Workers 临时工")
tab_exc       = _tab_refs.get("Exceptions 异常管理")
tab_dayoff    = _tab_refs.get("Day Off 调休")

# ============================================================
# HELPERS
# ============================================================
DAY_COLS = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
DAY_LABELS = ["Mon 周一", "Tue 周二", "Wed 周三", "Thu 周四", "Fri 周五", "Sat 周六", "Sun 周日"]

_ROW_EVEN = "background-color: #1e3a52; color: #ffffff"
_ROW_ODD  = "background-color: #162e42; color: #ffffff"

def style_table(df):
    def _zebra(x):
        colors = [_ROW_EVEN if i % 2 == 0 else _ROW_ODD for i in range(len(x))]
        return pd.DataFrame(
            [[c] * len(x.columns) for c in colors],
            index=x.index, columns=x.columns
        )
    return df.style.apply(_zebra, axis=None).set_properties(**{"white-space": "nowrap"})


def load_labor():
    with engine.connect() as conn:
        return pd.read_sql(text(f"SELECT * FROM {S}.dim_labor ORDER BY company, labor_id"), conn)


def load_shifts():
    with engine.connect() as conn:
        return pd.read_sql(text(f"SELECT * FROM {S}.dim_shift ORDER BY shift_name"), conn)


def load_schedules():
    with engine.connect() as conn:
        return pd.read_sql(text(f"""
            SELECT
                fs.schedule_id,
                l.labor_id,
                l.name,
                l.company,
                l.working_area,
                l.role,
                s.shift_name,
                fs.monday, fs.tuesday, fs.wednesday, fs.thursday,
                fs.friday, fs.saturday, fs.sunday
            FROM {S}.fact_schedule fs
            JOIN {S}.dim_labor l ON l.labor_id = fs.labor_id
            JOIN {S}.dim_shift  s ON s.shift_id  = fs.shift_id
            WHERE l.is_valid = TRUE
            ORDER BY l.company, l.labor_id
        """), conn)


# ============================================================
# TAB 1 — WORKERS
# ============================================================
if tab_workers is not None:
    with tab_workers:
        df_labor = load_labor()
        COMPANIES     = load_distinct("company")
        WORKING_AREAS = load_distinct("working_area")
        ROLES         = load_distinct("role")

        display_cols = ["labor_id", "name", "company", "working_area", "role", "is_valid", "created_at"]
        st.dataframe(style_table(df_labor[display_cols]), use_container_width=True, hide_index=True)
        st.divider()

        if st.session_state.get("workers_msg"):
            st.success(st.session_state.pop("workers_msg"))

        action = st.radio("Action", ["Add", "Edit"], horizontal=True, key="labor_action")

        if action == "Add":
            with st.form("add_labor"):
                c1, c2, c3 = st.columns(3)
                labor_id = c1.text_input("Labor ID *")
                name     = c2.text_input("Name *")
                company  = c3.selectbox("Company *", COMPANIES)
                if st.form_submit_button("Add Worker"):
                    if not labor_id or not name or not company:
                        st.error("Labor ID, Name, and Company are required.")
                    else:
                        try:
                            with engine.begin() as conn:
                                conn.execute(text(f"""
                                    INSERT INTO {S}.dim_labor (labor_id, name, company)
                                    VALUES (:labor_id, :name, :company)
                                """), {"labor_id": labor_id, "name": name, "company": company})
                            st.session_state["workers_msg"] = (
                                f"Worker added — **{name}** ({labor_id})  ·  Company: {company}"
                            )
                            st.rerun()
                        except Exception as e:
                            if "UniqueViolation" in type(e).__name__ or "duplicate key" in str(e).lower():
                                st.error(f"Worker **{labor_id}** already exists. Please check the Labor ID — this person may have been added before.")
                            else:
                                st.error(str(e))

        elif action == "Edit":
            if df_labor.empty:
                st.info("No workers yet.")
            else:
                options = {f"{r['name']}  —  {r['labor_id']}": r["labor_id"] for _, r in df_labor.iterrows()}
                sel_label = st.selectbox("Select worker", list(options.keys()), key="edit_labor_sel")
                sel = options[sel_label]
                row = df_labor[df_labor["labor_id"] == sel].iloc[0]
                with st.form("edit_labor"):
                    c1, c2, c3 = st.columns(3)
                    name     = c1.text_input("Name",  value=row["name"])
                    _co_val  = row["company"] if row["company"] in COMPANIES else ""
                    company  = c2.selectbox("Company", COMPANIES, index=COMPANIES.index(_co_val))
                    is_valid = c3.checkbox("Active",   value=bool(row["is_valid"]))
                    if st.form_submit_button("Save Changes"):
                        with engine.begin() as conn:
                            conn.execute(text(f"""
                                UPDATE {S}.dim_labor
                                SET name=:name, company=:company, is_valid=:is_valid
                                WHERE labor_id=:labor_id
                            """), {"name": name, "company": company, "is_valid": is_valid,
                                   "labor_id": sel})
                        st.session_state["workers_msg"] = (
                            f"Worker updated — **{name}** ({sel})\n\n"
                            f"Company: {company}  ·  {'Active' if is_valid else 'Inactive'}"
                        )
                        st.rerun()

                st.divider()
                st.markdown("**Schedule**")
                with engine.connect() as conn:
                    _ws = pd.read_sql(text(f"""
                        SELECT s.shift_name, s.clock_in, s.lunch_out, s.lunch_in, s.clock_out,
                               fs.monday, fs.tuesday, fs.wednesday, fs.thursday,
                               fs.friday, fs.saturday, fs.sunday
                        FROM {S}.fact_schedule fs
                        JOIN {S}.dim_shift s ON s.shift_id = fs.shift_id
                        WHERE fs.labor_id = :lid
                    """), conn, params={"lid": sel})
                if _ws.empty:
                    st.info("No schedule assigned.")
                else:
                    _r    = _ws.iloc[0]
                    _days = ", ".join([DAY_LABELS[i] for i, c in enumerate(DAY_COLS) if _r[c]]) or "None"
                    _wa_v = row["working_area"] if pd.notna(row["working_area"]) and row["working_area"] else "—"
                    _ro_v = row["role"]         if pd.notna(row["role"])         and row["role"]         else "—"
                    _sc1, _sc2, _sc3 = st.columns(3)
                    _sc1.text_input("Working Area", value=_wa_v,              disabled=True, key="ws_area")
                    _sc2.text_input("Role",         value=_ro_v,              disabled=True, key="ws_role")
                    _sc3.text_input("Shift",        value=_r["shift_name"],   disabled=True, key="ws_shift")
                    _sc4, _sc5, _sc6, _sc7 = st.columns(4)
                    _sc4.text_input("Working Days", value=_days,                                         disabled=True, key="ws_days")
                    _sc5.text_input("Clock In",     value=str(_r["clock_in"]),                           disabled=True, key="ws_ci")
                    _sc6.text_input("Lunch",        value=f"{_r['lunch_out']} – {_r['lunch_in']}",       disabled=True, key="ws_lunch")
                    _sc7.text_input("Clock Out",    value=str(_r["clock_out"]),                          disabled=True, key="ws_co")

                st.divider()
                _pending_del = st.session_state.get("confirm_del_worker")
                if _pending_del and _pending_del["id"] != sel:
                    st.session_state.pop("confirm_del_worker", None)
                    _pending_del = None
                if _pending_del:
                    st.warning(
                        f"Are you sure you want to permanently remove **{_pending_del['name']}** "
                        f"({_pending_del['id']}) from {_pending_del['company']}? "
                        f"Their schedule will also be deleted."
                    )
                    c_yes, c_no = st.columns(2)
                    if c_yes.button("Yes, delete", type="primary", key="del_labor_yes"):
                        try:
                            with engine.begin() as conn:
                                conn.execute(text(f"DELETE FROM {S}.fact_schedule WHERE labor_id=:id"), {"id": _pending_del["id"]})
                                conn.execute(text(f"DELETE FROM {S}.dim_labor    WHERE labor_id=:id"), {"id": _pending_del["id"]})
                            st.session_state.pop("confirm_del_worker", None)
                            st.session_state["workers_msg"] = (
                                f"Worker removed — **{_pending_del['name']}** ({_pending_del['id']})  ·  {_pending_del['company']}"
                            )
                            st.rerun()
                        except Exception as e:
                            st.error(str(e))
                    if c_no.button("Cancel", key="del_labor_cancel"):
                        st.session_state.pop("confirm_del_worker", None)
                        st.rerun()
                else:
                    st.markdown('<span class="del-marker"></span>', unsafe_allow_html=True)
                    _, _del_col = st.columns([3, 1])
                    with _del_col:
                        if st.button("Delete Worker", key="del_labor_btn", use_container_width=True):
                            st.session_state["confirm_del_worker"] = {
                                "id": sel, "name": row["name"], "company": row["company"]
                            }
                            st.rerun()


# ============================================================
# TAB 2 — SHIFTS
# ============================================================
if tab_shifts is not None:
    with tab_shifts:
        df_shifts = load_shifts()

        display_cols = ["shift_id", "shift_name", "clock_in", "lunch_out", "lunch_in", "clock_out"]
        st.dataframe(style_table(df_shifts[display_cols]), use_container_width=True, hide_index=True)
        st.divider()

        if st.session_state.get("shifts_msg"):
            st.success(st.session_state.pop("shifts_msg"))

        action = st.radio("Action", ["Add", "Edit", "Delete"], horizontal=True, key="shift_action")

        if action == "Add":
            with st.form("add_shift"):
                shift_name = st.text_input("Shift Name *")
                c1, c2, c3, c4 = st.columns(4)
                clock_in  = c1.time_input("Clock In")
                lunch_out = c2.time_input("Lunch Out")
                lunch_in  = c3.time_input("Lunch In")
                clock_out = c4.time_input("Clock Out")
                if st.form_submit_button("Add Shift"):
                    if not shift_name:
                        st.error("Shift name is required.")
                    else:
                        try:
                            with engine.begin() as conn:
                                conn.execute(text(f"""
                                    INSERT INTO {S}.dim_shift (shift_name, clock_in, lunch_out, lunch_in, clock_out)
                                    VALUES (:name, :ci, :lo, :li, :co)
                                """), {"name": shift_name, "ci": str(clock_in), "lo": str(lunch_out),
                                      "li": str(lunch_in), "co": str(clock_out)})
                            st.session_state["shifts_msg"] = (
                                f"Shift added — **{shift_name}**\n\n"
                                f"Clock In: {clock_in}  ·  Lunch: {lunch_out} – {lunch_in}  ·  Clock Out: {clock_out}"
                            )
                            st.rerun()
                        except Exception as e:
                            st.error(str(e))

        elif action == "Edit":
            if df_shifts.empty:
                st.info("No shifts yet.")
            else:
                options = {f"{r['shift_name']} (ID {r['shift_id']})": r["shift_id"]
                           for _, r in df_shifts.iterrows()}
                sel_label = st.selectbox("Select shift", list(options.keys()), key="edit_shift_sel")
                sel_id    = options[sel_label]
                row       = df_shifts[df_shifts["shift_id"] == sel_id].iloc[0]
                with st.form("edit_shift"):
                    shift_name = st.text_input("Shift Name", value=row["shift_name"])
                    c1, c2, c3, c4 = st.columns(4)
                    clock_in  = c1.time_input("Clock In",  value=row["clock_in"])
                    lunch_out = c2.time_input("Lunch Out", value=row["lunch_out"])
                    lunch_in  = c3.time_input("Lunch In",  value=row["lunch_in"])
                    clock_out = c4.time_input("Clock Out", value=row["clock_out"])
                    if st.form_submit_button("Save Changes"):
                        with engine.begin() as conn:
                            conn.execute(text(f"""
                                UPDATE {S}.dim_shift
                                SET shift_name=:name, clock_in=:ci, lunch_out=:lo, lunch_in=:li, clock_out=:co
                                WHERE shift_id=:id
                            """), {"name": shift_name, "ci": str(clock_in), "lo": str(lunch_out),
                                   "li": str(lunch_in), "co": str(clock_out), "id": sel_id})
                        st.session_state["shifts_msg"] = (
                            f"Shift updated — **{shift_name}**\n\n"
                            f"Clock In: {clock_in}  ·  Lunch: {lunch_out} – {lunch_in}  ·  Clock Out: {clock_out}"
                        )
                        st.rerun()

        elif action == "Delete":
            if df_shifts.empty:
                st.info("No shifts yet.")
            else:
                options  = {f"{r['shift_name']} (ID {r['shift_id']})": r["shift_id"]
                            for _, r in df_shifts.iterrows()}
                sel_label = st.selectbox("Select shift to delete", list(options.keys()), key="del_shift_sel")
                sel_id    = options[sel_label]
                sel_name  = df_shifts[df_shifts["shift_id"] == sel_id].iloc[0]["shift_name"]

                if st.button("Delete Shift", type="primary", key="del_shift_btn"):
                    st.session_state["confirm_del_shift"] = {"id": sel_id, "name": sel_name}

                pending = st.session_state.get("confirm_del_shift")
                if pending:
                    st.warning(
                        f"Are you sure you want to delete shift **{pending['name']}**? "
                        f"This will fail if workers are still assigned to it."
                    )
                    c_yes, c_no = st.columns(2)
                    if c_yes.button("Yes, delete", type="primary", key="del_shift_yes"):
                        try:
                            with engine.begin() as conn:
                                conn.execute(text(f"DELETE FROM {S}.dim_shift WHERE shift_id=:id"), {"id": pending["id"]})
                            st.session_state.pop("confirm_del_shift", None)
                            st.session_state["shifts_msg"] = f"Shift removed — **{pending['name']}**"
                            st.rerun()
                        except Exception as e:
                            st.error(str(e))
                    if c_no.button("Cancel", key="del_shift_cancel"):
                        st.session_state.pop("confirm_del_shift", None)
                        st.rerun()


# ============================================================
# TAB 3 — SCHEDULES
# ============================================================
if tab_schedules is not None:
    with tab_schedules:
        df_schedules = load_schedules()
        df_labor     = load_labor()
        df_shifts    = load_shifts()

        st.dataframe(style_table(df_schedules.drop(columns=["schedule_id"])), use_container_width=True, hide_index=True)
        st.divider()

        if st.session_state.get("schedules_msg"):
            st.success(st.session_state.pop("schedules_msg"))

        _SHORT_DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

        if df_shifts.empty:
            st.info("Create at least one shift first.")
        elif df_labor.empty:
            st.info("No active workers yet.")
        else:
            labor_opts  = {f"{r['name']} ({r['labor_id']})": r["labor_id"]
                           for _, r in df_labor.iterrows()}
            shift_opts  = {f"{r['shift_name']} (ID {r['shift_id']})": r["shift_id"]
                           for _, r in df_shifts.iterrows()}
            shift_keys  = list(shift_opts.keys())

            sel_labor_label = st.selectbox("Worker", list(labor_opts.keys()), key="upsert_sched_labor")
            sel_labor_id    = labor_opts[sel_labor_label]

            _wr = df_labor[df_labor["labor_id"] == sel_labor_id].iloc[0]
            _wi1, _wi2 = st.columns(2)
            _wi1.text_input("Company", value=_wr["company"],                               disabled=True, key="wi_company")
            _wi2.text_input("Status",  value="Active" if _wr["is_valid"] else "Inactive",  disabled=True, key="wi_status")

            with engine.connect() as conn:
                existing = pd.read_sql(
                    text(f"SELECT * FROM {S}.fact_schedule WHERE labor_id=:lid"),
                    conn, params={"lid": sel_labor_id}
                )
            has_schedule = not existing.empty

            if has_schedule:
                ex           = existing.iloc[0]
                cur_shift_k  = next((k for k, v in shift_opts.items() if v == ex["shift_id"]), shift_keys[0])
                shift_idx    = shift_keys.index(cur_shift_k)
                default_days = [bool(ex[DAY_COLS[i]]) for i in range(7)]
            else:
                shift_idx    = 0
                default_days = [False] * 7

            _sched_areas = [""] + sorted([v for v in df_labor["working_area"].dropna().unique() if v != ""])
            _sched_roles = [""] + sorted([v for v in df_labor["role"].dropna().unique() if v != ""])
            _cur_wa = _wr["working_area"] if pd.notna(_wr["working_area"]) and _wr["working_area"] in _sched_areas else ""
            _cur_ro = _wr["role"]         if pd.notna(_wr["role"])         and _wr["role"]         in _sched_roles else ""

            with st.form("upsert_schedule"):
                _wa_col, _ro_col = st.columns(2)
                working_area = _wa_col.selectbox("Working Area", _sched_areas,
                                                 index=_sched_areas.index(_cur_wa))
                role         = _ro_col.selectbox("Role", _sched_roles,
                                                 index=_sched_roles.index(_cur_ro))
                sel_shift = st.selectbox("Shift", shift_keys, index=shift_idx)
                st.write("Working days:")
                day_cols_ui = st.columns(7)
                day_vals = [day_cols_ui[i].checkbox(DAY_LABELS[i], value=default_days[i],
                                                    key=f"upsert_day_{i}") for i in range(7)]
                btn_label = "Save Changes" if has_schedule else "Add Schedule"
                if st.form_submit_button(btn_label):
                    params = {
                        "lid": sel_labor_id, "sid": shift_opts[sel_shift],
                        "mon": day_vals[0], "tue": day_vals[1], "wed": day_vals[2],
                        "thu": day_vals[3], "fri": day_vals[4], "sat": day_vals[5],
                        "sun": day_vals[6],
                    }
                    with engine.begin() as conn:
                        conn.execute(text(f"""
                            UPDATE {S}.dim_labor
                            SET working_area=:working_area, role=:role
                            WHERE labor_id=:lid
                        """), {"working_area": working_area or None, "role": role or None,
                               "lid": sel_labor_id})
                        if has_schedule:
                            conn.execute(text(f"""
                                UPDATE {S}.fact_schedule
                                SET shift_id=:sid,
                                    monday=:mon, tuesday=:tue, wednesday=:wed, thursday=:thu,
                                    friday=:fri, saturday=:sat, sunday=:sun
                                WHERE labor_id=:lid
                            """), params)
                            msg_prefix = "Schedule updated"
                        else:
                            conn.execute(text(f"""
                                INSERT INTO {S}.fact_schedule
                                    (labor_id, shift_id, monday, tuesday, wednesday, thursday, friday, saturday, sunday)
                                VALUES (:lid, :sid, :mon, :tue, :wed, :thu, :fri, :sat, :sun)
                            """), params)
                            msg_prefix = "Schedule added"
                    active_days = ", ".join([_SHORT_DAYS[i] for i, v in enumerate(day_vals) if v]) or "None"
                    shift_name  = sel_shift.split(" (ID")[0]
                    st.session_state["schedules_msg"] = (
                        f"{msg_prefix} — **{sel_labor_label}**\n\n"
                        f"Area: {working_area or '—'}  ·  Role: {role or '—'}  ·  "
                        f"Shift: {shift_name}  ·  Days: {active_days}"
                    )
                    st.rerun()

            if has_schedule:
                st.divider()
                _sched_id = int(existing.iloc[0]["schedule_id"])
                _pending_sched = st.session_state.get("confirm_del_sched")
                if _pending_sched and _pending_sched["id"] != _sched_id:
                    st.session_state.pop("confirm_del_sched", None)
                    _pending_sched = None
                if _pending_sched:
                    st.warning(
                        f"Are you sure you want to remove the entire schedule for **{sel_labor_label}**? "
                        f"This cannot be undone."
                    )
                    c_yes, c_no = st.columns(2)
                    if c_yes.button("Yes, delete", type="primary", key="del_sched_yes"):
                        with engine.begin() as conn:
                            conn.execute(text(f"DELETE FROM {S}.fact_schedule WHERE schedule_id=:id"), {"id": _sched_id})
                            conn.execute(text(f"""
                                UPDATE {S}.dim_labor
                                SET working_area=NULL, role=NULL
                                WHERE labor_id=:lid
                            """), {"lid": sel_labor_id})
                        st.session_state.pop("confirm_del_sched", None)
                        st.session_state["schedules_msg"] = f"Schedule removed — **{sel_labor_label}**"
                        st.rerun()
                    if c_no.button("Cancel", key="del_sched_cancel"):
                        st.session_state.pop("confirm_del_sched", None)
                        st.rerun()
                else:
                    st.markdown('<span class="del-marker"></span>', unsafe_allow_html=True)
                    _, _del_col = st.columns([3, 1])
                    with _del_col:
                        if st.button("Delete Schedule", key="del_sched_btn", use_container_width=True):
                            st.session_state["confirm_del_sched"] = {"id": _sched_id, "label": sel_labor_label}
                            st.rerun()


# ============================================================
# TAB 4 — AREA ANALYSIS
# ============================================================
if tab_area is not None:
    with tab_area:
        st.subheader("Area Analysis 区域分析")

        _AA_DAY_COLS   = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
        _AA_DAY_LABELS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

        with engine.connect() as conn:
            _aa_areas = [r[0] for r in conn.execute(text(f"""
                SELECT DISTINCT working_area FROM {S}.dim_labor
                WHERE working_area IS NOT NULL AND working_area != '' AND is_valid = TRUE
                ORDER BY working_area
            """)).fetchall()]

        if not _aa_areas:
            st.info("No working areas found in active workers.")
        else:
            sel_area = st.selectbox("Working Area", _aa_areas, key="area_sel")

            # ── Currently scheduled counts per day for this area ──
            with engine.connect() as conn:
                _aa_sched_row = conn.execute(text(f"""
                    SELECT
                        COALESCE(SUM(CASE WHEN fs.monday    THEN 1 ELSE 0 END), 0),
                        COALESCE(SUM(CASE WHEN fs.tuesday   THEN 1 ELSE 0 END), 0),
                        COALESCE(SUM(CASE WHEN fs.wednesday THEN 1 ELSE 0 END), 0),
                        COALESCE(SUM(CASE WHEN fs.thursday  THEN 1 ELSE 0 END), 0),
                        COALESCE(SUM(CASE WHEN fs.friday    THEN 1 ELSE 0 END), 0),
                        COALESCE(SUM(CASE WHEN fs.saturday  THEN 1 ELSE 0 END), 0),
                        COALESCE(SUM(CASE WHEN fs.sunday    THEN 1 ELSE 0 END), 0)
                    FROM {S}.fact_schedule fs
                    JOIN {S}.dim_labor l ON l.labor_id = fs.labor_id
                    WHERE l.working_area = :area AND l.is_valid = TRUE
                """), {"area": sel_area}).fetchone()
            scheduled = [int(_aa_sched_row[i]) for i in range(7)]

            # ── Target for this area ──────────────────────────────
            with engine.connect() as conn:
                _aa_tgt_row = conn.execute(text(f"""
                    SELECT monday, tuesday, wednesday, thursday, friday, saturday, sunday
                    FROM {S}.area_targets WHERE working_area = :area
                """), {"area": sel_area}).fetchone()
            target = [int(_aa_tgt_row[i]) for i in range(7)] if _aa_tgt_row else [0] * 7

            missing = [target[i] - scheduled[i] for i in range(7)]

            # ── Styled 3-row grid ─────────────────────────────────
            _aa_grid = pd.DataFrame(
                [scheduled, target, missing],
                index=["Currently Scheduled", "Target", "Missing"],
                columns=_AA_DAY_LABELS,
            )

            def _style_area_grid(df):
                styles = pd.DataFrame("", index=df.index, columns=df.columns)
                styles.loc["Currently Scheduled"] = "background-color: #1e3a52; color: #ffffff"
                styles.loc["Target"]              = "background-color: #162e42; color: #ffffff"
                for col in df.columns:
                    v = df.loc["Missing", col]
                    if v > 0:
                        styles.loc["Missing", col] = "background-color: #c62828; color: #ffffff; font-weight: bold"
                    elif v < 0:
                        styles.loc["Missing", col] = "background-color: #1565c0; color: #ffffff"
                    else:
                        styles.loc["Missing", col] = "background-color: #2e7d32; color: #ffffff"
                return styles

            st.dataframe(
                _aa_grid.style.apply(_style_area_grid, axis=None),
                use_container_width=True,
            )
            st.caption("Missing = Target − Scheduled  ·  🔴 understaffed  ·  🟢 exact  ·  🔵 overstaffed")

            st.divider()

            # ── Set Target form ───────────────────────────────────
            st.markdown(f"**Set Targets for {sel_area}**")
            with st.form("area_target_form"):
                _aa_tcols = st.columns(7)
                _aa_inputs = [
                    _aa_tcols[i].number_input(
                        _AA_DAY_LABELS[i], min_value=0, value=target[i], step=1, key=f"aa_tgt_{i}"
                    )
                    for i in range(7)
                ]
                if st.form_submit_button("Save Targets", use_container_width=True):
                    with engine.begin() as conn:
                        conn.execute(text(f"""
                            INSERT INTO {S}.area_targets
                                (working_area, monday, tuesday, wednesday, thursday, friday, saturday, sunday)
                            VALUES (:area, :mon, :tue, :wed, :thu, :fri, :sat, :sun)
                            ON CONFLICT (working_area) DO UPDATE SET
                                monday=:mon, tuesday=:tue, wednesday=:wed, thursday=:thu,
                                friday=:fri, saturday=:sat, sunday=:sun, updated_at=NOW()
                        """), {
                            "area": sel_area,
                            "mon": int(_aa_inputs[0]), "tue": int(_aa_inputs[1]),
                            "wed": int(_aa_inputs[2]), "thu": int(_aa_inputs[3]),
                            "fri": int(_aa_inputs[4]), "sat": int(_aa_inputs[5]),
                            "sun": int(_aa_inputs[6]),
                        })
                    st.success(f"Targets saved for **{sel_area}**.")
                    st.rerun()

            st.divider()

            # ── All-areas missing overview ────────────────────────
            st.markdown("**All Areas — Current Schedule Overview**")

            with engine.connect() as conn:
                _aa_all_sched = pd.read_sql(text(f"""
                    SELECT
                        l.working_area,
                        COALESCE(SUM(CASE WHEN fs.monday    THEN 1 ELSE 0 END), 0) AS monday,
                        COALESCE(SUM(CASE WHEN fs.tuesday   THEN 1 ELSE 0 END), 0) AS tuesday,
                        COALESCE(SUM(CASE WHEN fs.wednesday THEN 1 ELSE 0 END), 0) AS wednesday,
                        COALESCE(SUM(CASE WHEN fs.thursday  THEN 1 ELSE 0 END), 0) AS thursday,
                        COALESCE(SUM(CASE WHEN fs.friday    THEN 1 ELSE 0 END), 0) AS friday,
                        COALESCE(SUM(CASE WHEN fs.saturday  THEN 1 ELSE 0 END), 0) AS saturday,
                        COALESCE(SUM(CASE WHEN fs.sunday    THEN 1 ELSE 0 END), 0) AS sunday
                    FROM {S}.fact_schedule fs
                    JOIN {S}.dim_labor l ON l.labor_id = fs.labor_id
                    WHERE l.is_valid = TRUE AND l.working_area IS NOT NULL AND l.working_area != ''
                    GROUP BY l.working_area
                    ORDER BY l.working_area
                """), conn)

            if _aa_all_sched.empty:
                st.info("No scheduled workers with assigned areas yet.")
            else:
                _aa_overview = (
                    _aa_all_sched
                    .rename(columns={dc: dl for dc, dl in zip(_AA_DAY_COLS, _AA_DAY_LABELS)})
                    .rename(columns={"working_area": "Area"})
                    .set_index("Area")
                )
                st.dataframe(style_table(_aa_overview), use_container_width=True)


# ============================================================
# TAB 5 — TODAY'S EXPORT
# ============================================================
if tab_export is not None:
    with tab_export:
        from datetime import date
        from openpyxl import Workbook
        from openpyxl.styles import Font, Border, Side, Alignment

        _side   = Side(style="thin")
        _border = Border(left=_side, right=_side, top=_side, bottom=_side)
        _FONT   = "Times New Roman"
        _SZ     = 30

        st.subheader("Attendance Export 签到表导出")

        _DAY_COLS = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
        chosen_date = st.date_input("Select date", value=date.today())
        day_col     = _DAY_COLS[chosen_date.weekday()]  # weekday(): Mon=0 … Sun=6

        query = text(f"""
            SELECT
                l.company                        AS "劳务公司",
                l.labor_id                       AS "劳务ID",
                l.name                           AS "劳务名字",
                l.role                           AS "Role",
                l.working_area                   AS "Area",
                s.shift_name                     AS "Shift",
                ''::text AS "Clock In",
                ''::text AS "Lunch Out",
                ''::text AS "Lunch In",
                ''::text AS "Clock Out"
            FROM {S}.fact_schedule fs
            JOIN {S}.dim_labor l ON l.labor_id = fs.labor_id
            JOIN {S}.dim_shift  s ON s.shift_id  = fs.shift_id
            WHERE l.is_valid = TRUE
              AND (
                    (
                        fs.{day_col} = TRUE
                        AND NOT EXISTS (
                            SELECT 1 FROM {S}.fact_day_off_change doc
                            WHERE doc.labor_id = fs.labor_id
                              AND doc.change_date = :chosen_date
                              AND doc.new_status = 'off'
                        )
                    )
                    OR EXISTS (
                        SELECT 1 FROM {S}.fact_day_off_change doc
                        WHERE doc.labor_id = fs.labor_id
                          AND doc.change_date = :chosen_date
                          AND doc.new_status = 'work'
                    )
              )
            ORDER BY l.company, s.shift_name, l.role, l.name
        """)

        with engine.connect() as conn:
            df_export = pd.read_sql(query, conn, params={"chosen_date": str(chosen_date)})

        if df_export.empty:
            st.info(f"No workers scheduled for {chosen_date.strftime('%A, %Y-%m-%d')}.")
        else:
            st.dataframe(style_table(df_export), use_container_width=True, hide_index=True)
            st.caption(f"{len(df_export)} workers scheduled on {chosen_date.strftime('%A, %Y-%m-%d')}.")

            export_cols = ["劳务ID", "劳务名字", "Shift", "Role", "Area",
                           "Clock In", "Lunch Out", "Lunch In", "Clock Out"]
            lead_roles  = {"Lead", "Assistant Lead"}

            from openpyxl.utils import get_column_letter

            _ROW_H   = 70    # 93px ≈ 70pt
            _TITLE_SZ = 60
            _n_cols  = len(export_cols)

            buf = io.BytesIO()
            wb  = Workbook()
            wb.remove(wb.active)

            for company, group in df_export.groupby("劳务公司"):
                ws = wb.create_sheet(title=str(company)[:31])

                # Row 1: agency name — merged, centered, size 60
                ws.merge_cells(start_row=1, start_column=1,
                               end_row=1,   end_column=_n_cols)
                for col_idx in range(1, _n_cols + 1):
                    cell        = ws.cell(row=1, column=col_idx)
                    cell.border = _border
                    cell.font   = Font(name=_FONT, bold=True, size=_TITLE_SZ)
                title_cell           = ws.cell(row=1, column=1)
                title_cell.value     = str(company)
                title_cell.alignment = Alignment(horizontal="center", vertical="center")

                # Row 2: column headers
                for col_idx, col_name in enumerate(export_cols, 1):
                    cell        = ws.cell(row=2, column=col_idx, value=col_name)
                    cell.font   = Font(name=_FONT, bold=True, size=_SZ)
                    cell.border = _border

                # Row 3+: data rows
                for row_offset, (_, r) in enumerate(group[export_cols].iterrows(), 3):
                    is_lead = r["Role"] in lead_roles
                    for col_idx, col_name in enumerate(export_cols, 1):
                        value = "AL" if col_name == "Role" and r[col_name] == "Assistant Lead" else r[col_name]
                        cell        = ws.cell(row=row_offset, column=col_idx, value=value)
                        cell.border = _border
                        cell.font   = Font(name=_FONT, bold=is_lead, italic=is_lead, size=_SZ)

                # Set row height (93px ≈ 70pt) for every used row
                total_rows = 2 + len(group)
                for row_idx in range(1, total_rows + 1):
                    ws.row_dimensions[row_idx].height = _ROW_H

                # Hide unused columns beyond the data range
                for extra in range(_n_cols + 1, _n_cols + 50):
                    ws.column_dimensions[get_column_letter(extra)].hidden = True

            wb.save(buf)
            buf.seek(0)

            st.download_button(
                label=f"Download Excel — {chosen_date.strftime('%Y-%m-%d')}",
                data=buf,
                file_name=f"attendance_{chosen_date}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )


# ============================================================
# TAB 6 — TEMP WORKERS
# ============================================================
if tab_temp is not None:
    with tab_temp:
        from datetime import date as _date, timedelta
        from openpyxl import Workbook as _Workbook
        from openpyxl.styles import Font as _Font, Border as _Border, Side as _Side

        _t_side   = _Side(style="thin")
        _t_border = _Border(left=_t_side, right=_t_side, top=_t_side, bottom=_t_side)
        _T_FONT   = "Times New Roman"
        _T_SZ     = 30

        st.subheader("Temporary Workers 临时工")

        def load_temp(work_date):
            with engine.connect() as conn:
                return pd.read_sql(text(f"""
                    SELECT t.temp_id, l.labor_id, l.name, l.company, l.role
                    FROM {S}.fact_temp_schedule t
                    JOIN {S}.dim_labor l ON l.labor_id = t.labor_id
                    WHERE t.work_date = :d
                      AND l.is_valid = TRUE
                    ORDER BY l.company, l.name
                """), conn, params={"d": str(work_date)})

        temp_date = st.date_input("Date", value=_date.today() + timedelta(days=1), key="temp_date")
        df_temp   = load_temp(temp_date)

        if df_temp.empty:
            st.info(f"No temp workers for {temp_date.strftime('%A, %Y-%m-%d')}.")
        else:
            st.dataframe(style_table(df_temp[["labor_id", "name", "company", "role"]]),
                         use_container_width=True, hide_index=True)
            st.caption(f"{len(df_temp)} temp workers on {temp_date.strftime('%A, %Y-%m-%d')}.")

        st.divider()
        t_action = st.radio("Action", ["Add", "Remove"], horizontal=True, key="temp_action")

        if t_action == "Add":
            t_search   = st.text_input("Search by name", key="temp_search", placeholder="Type to filter...")
            df_all     = load_labor()
            existing   = set(df_temp["labor_id"].tolist()) if not df_temp.empty else set()
            available  = df_all[df_all["is_valid"] == True][~df_all["labor_id"].isin(existing)]
            t_filtered = (available[available["name"].str.contains(t_search, case=False, na=False)]
                          if t_search else available)

            if t_filtered.empty:
                st.info("No matching workers available.")
            else:
                t_opts  = {f"{r['name']}  —  {r['labor_id']}  ({r['company']})": r["labor_id"]
                           for _, r in t_filtered.iterrows()}
                t_label = st.selectbox("Select worker", list(t_opts.keys()), key="temp_sel")
                if st.button("Add to schedule", key="temp_add_btn"):
                    with engine.begin() as conn:
                        conn.execute(text(f"""
                            INSERT INTO {S}.fact_temp_schedule (labor_id, work_date)
                            VALUES (:lid, :d)
                            ON CONFLICT (labor_id, work_date) DO NOTHING
                        """), {"lid": t_opts[t_label], "d": str(temp_date)})
                    st.success("Added.")
                    st.rerun()

        elif t_action == "Remove":
            if df_temp.empty:
                st.info("No temp workers to remove.")
            else:
                rem_opts  = {f"{r['name']}  —  {r['labor_id']}": r["temp_id"]
                             for _, r in df_temp.iterrows()}
                rem_label = st.selectbox("Select worker to remove", list(rem_opts.keys()), key="temp_rem_sel")
                if st.button("Remove", type="primary", key="temp_rem_btn"):
                    with engine.begin() as conn:
                        conn.execute(text(f"DELETE FROM {S}.fact_temp_schedule WHERE temp_id=:id"),
                                     {"id": rem_opts[rem_label]})
                    st.success("Removed.")
                    st.rerun()

        if not df_temp.empty:
            st.divider()
            t_export_cols = ["劳务公司", "劳务ID", "劳务名字", "Role",
                             "Clock In", "Lunch Out", "Lunch In", "Clock Out"]
            t_lead_roles  = {"Lead", "Assistant Lead"}

            df_temp_xl = df_temp.sort_values("company").copy()
            df_temp_xl = df_temp_xl.rename(columns={
                "company":  "劳务公司",
                "labor_id": "劳务ID",
                "name":     "劳务名字",
                "role":     "Role",
            })
            df_temp_xl["Clock In"]  = ""
            df_temp_xl["Lunch Out"] = ""
            df_temp_xl["Lunch In"]  = ""
            df_temp_xl["Clock Out"] = ""

            t_buf = io.BytesIO()
            t_wb  = _Workbook()
            ws    = t_wb.active
            ws.title = "Temp Attendance"

            for col_idx, col_name in enumerate(t_export_cols, 1):
                cell        = ws.cell(row=1, column=col_idx, value=col_name)
                cell.font   = _Font(name=_T_FONT, bold=True, size=_T_SZ)
                cell.border = _t_border

            for row_offset, (_, r) in enumerate(df_temp_xl[t_export_cols].iterrows(), 2):
                is_lead = r["Role"] in t_lead_roles
                for col_idx, col_name in enumerate(t_export_cols, 1):
                    cell        = ws.cell(row=row_offset, column=col_idx, value=r[col_name])
                    cell.border = _t_border
                    cell.font   = _Font(name=_T_FONT, bold=is_lead, italic=is_lead, size=_T_SZ)

            t_wb.save(t_buf)
            t_buf.seek(0)

            st.download_button(
                label=f"Download Excel — {temp_date.strftime('%Y-%m-%d')}",
                data=t_buf,
                file_name=f"temp_attendance_{temp_date}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )


# ============================================================
# TAB 7 — EXCEPTIONS MANAGEMENT
# ============================================================
if tab_exc is not None:
    with tab_exc:
        from datetime import date as _edate

        st.subheader("Exceptions Management 异常管理")

        _REASON_LABELS = {"leave_early": "Leave Early", "move_day_off": "Move Day Off", "take_leave": "Take Leave"}
        _LEAVE_LABELS  = {"sick": "Sick", "not_reported": "Not Reported", "personal": "Personal"}

        def load_exceptions():
            with engine.connect() as conn:
                return pd.read_sql(text(f"""
                    SELECT
                        e.exception_id,
                        l.name          AS worker_name,
                        l.labor_id,
                        e.occur_date,
                        e.reason,
                        e.leave_early_type,
                        e.leave_early_time,
                        e.leave_date,
                        e.move_from_date,
                        e.move_to_date,
                        e.comment,
                        e.created_at
                    FROM {S}.fact_exceptions e
                    JOIN {S}.dim_labor l ON l.labor_id = e.labor_id
                    ORDER BY e.occur_date DESC, l.name
                """), conn)

        df_exc = load_exceptions()

        # ── Exceptions table with date filter ───────────────────
        fc1, fc2 = st.columns(2)
        exc_filter_date = fc1.date_input("Filter by Date Occurred", value=None, key="exc_filter_date")

        if not df_exc.empty:
            display_exc = df_exc.copy()
            if exc_filter_date:
                display_exc = display_exc[pd.to_datetime(display_exc["occur_date"]).dt.date == exc_filter_date]
            display_exc["reason"]           = display_exc["reason"].map(_REASON_LABELS)
            display_exc["leave_early_type"] = display_exc["leave_early_type"].map(_LEAVE_LABELS)
            if display_exc.empty:
                st.info(f"No exceptions on {exc_filter_date}.")
            else:
                st.dataframe(
                    style_table(display_exc[["worker_name", "labor_id", "occur_date", "reason",
                                             "leave_early_type", "leave_early_time", "leave_date",
                                             "move_from_date", "move_to_date", "comment"]]),
                    use_container_width=True, hide_index=True
                )
                st.caption(f"{len(display_exc)} exception(s){'  ·  filtered' if exc_filter_date else ''}.")
        else:
            st.info("No exceptions recorded yet.")

        st.divider()

        if st.session_state.get("exc_msg"):
            st.success(st.session_state.pop("exc_msg"))

        exc_action = st.radio("Action", ["Add", "Delete"], horizontal=True, key="exc_action")

        if exc_action == "Add":
            df_exc_labor = load_labor()
            active_labor = df_exc_labor[df_exc_labor["is_valid"] == True]

            if active_labor.empty:
                st.info("No active workers found.")
            else:
                exc_worker_opts = {
                    f"{r['name']} ({r['labor_id']})": r["labor_id"]
                    for _, r in active_labor.iterrows()
                }

                c1, c2 = st.columns(2)
                exc_occur_date     = c1.date_input("Date Occurred", value=_edate.today(), key="exc_occur_date")
                exc_worker_label   = c2.selectbox("Worker", list(exc_worker_opts.keys()), key="exc_worker_sel")
                exc_worker_id      = exc_worker_opts[exc_worker_label]

                exc_reason = st.selectbox("Reason", ["Leave Early", "Take Leave"], key="exc_reason_sel")

                exc_leave_type  = None
                exc_comment     = None
                exc_leave_date  = None

                exc_leave_time = None

                if exc_reason == "Leave Early":
                    exc_leave_type = st.selectbox("Leave Type", ["Sick", "Not Reported", "Personal"], key="exc_leave_type_sel")
                    exc_leave_time = st.time_input("Time Left", key="exc_leave_time")
                    if exc_leave_type == "Personal":
                        exc_comment = st.text_area("Comment", key="exc_comment", placeholder="Briefly describe what happened...")

                elif exc_reason == "Take Leave":
                    exc_leave_date = st.date_input("Leave Date", value=_edate.today(), key="exc_leave_date")

                if st.button("Record Exception", type="primary", key="exc_add_btn"):
                    try:
                        _reason_map = {"Leave Early": "leave_early", "Take Leave": "take_leave"}
                        _reason_val = _reason_map[exc_reason]
                        _leave_map  = {"Sick": "sick", "Not Reported": "not_reported", "Personal": "personal"}
                        _leave_val  = _leave_map.get(exc_leave_type) if exc_leave_type else None
                        with engine.begin() as conn:
                            conn.execute(text(f"""
                                INSERT INTO {S}.fact_exceptions
                                    (labor_id, occur_date, reason, leave_early_type, leave_early_time, leave_date, comment)
                                VALUES (:lid, :od, :reason, :let, :letm, :ld, :comment)
                            """), {
                                "lid":     exc_worker_id,
                                "od":      str(exc_occur_date),
                                "reason":  _reason_val,
                                "let":     _leave_val,
                                "letm":    str(exc_leave_time) if exc_leave_time else None,
                                "ld":      str(exc_leave_date) if exc_leave_date else None,
                                "comment": exc_comment or None,
                            })
                        st.session_state["exc_msg"] = (
                            f"Exception recorded — **{exc_worker_label}**  ·  "
                            f"{exc_occur_date}  ·  {exc_reason}"
                        )
                        st.rerun()
                    except Exception as e:
                        st.error(str(e))

        elif exc_action == "Delete":
            if df_exc.empty:
                st.info("No exceptions to delete.")
            else:
                del_exc_opts = {
                    f"{r['worker_name']} ({r['labor_id']})  —  {r['occur_date']}  —  {_REASON_LABELS.get(r['reason'], r['reason'])}": r["exception_id"]
                    for _, r in df_exc.iterrows()
                }
                del_exc_label = st.selectbox("Select exception", list(del_exc_opts.keys()), key="del_exc_sel")
                del_exc_id    = del_exc_opts[del_exc_label]

                if st.button("Delete Exception", type="primary", key="del_exc_btn"):
                    st.session_state["confirm_del_exc"] = {"id": del_exc_id, "label": del_exc_label}

                pending_exc = st.session_state.get("confirm_del_exc")
                if pending_exc:
                    st.warning(f"Delete exception: **{pending_exc['label']}**?")
                    cy, cn = st.columns(2)
                    if cy.button("Yes, delete", type="primary", key="del_exc_yes"):
                        with engine.begin() as conn:
                            conn.execute(text(f"DELETE FROM {S}.fact_exceptions WHERE exception_id=:id"),
                                         {"id": pending_exc["id"]})
                        st.session_state.pop("confirm_del_exc", None)
                        st.session_state["exc_msg"] = "Exception deleted."
                        st.rerun()
                    if cn.button("Cancel", key="del_exc_cancel"):
                        st.session_state.pop("confirm_del_exc", None)
                        st.rerun()



# ============================================================
# TAB 8 — DAY OFF CHANGES
# ============================================================
if tab_dayoff is not None:
    with tab_dayoff:
        from datetime import date as _ddate, timedelta as _dtimedelta

        st.subheader("Day Off 调休")

        df_do_labor  = load_labor()
        active_labor = df_do_labor[df_do_labor["is_valid"] == True]

        if active_labor.empty:
            st.info("No active workers found.")
        else:
            do_worker_opts  = {f"{r['name']} ({r['labor_id']})": r["labor_id"]
                               for _, r in active_labor.iterrows()}
            do_worker_label = st.selectbox("Worker", list(do_worker_opts.keys()), key="dayoff_worker_sel")
            do_labor_id     = do_worker_opts[do_worker_label]

            with engine.connect() as conn:
                do_sched = pd.read_sql(text(f"""
                    SELECT monday, tuesday, wednesday, thursday, friday, saturday, sunday
                    FROM {S}.fact_schedule WHERE labor_id = :lid
                """), conn, params={"lid": do_labor_id})

            if do_sched.empty:
                st.warning("This worker has no regular schedule on file — nothing to base a calendar on.")
            else:
                _do_row     = do_sched.iloc[0]
                _do_default = {i: ("work" if _do_row[DAY_COLS[i]] else "off") for i in range(7)}

                _do_dates = [_ddate.today() + _dtimedelta(days=i) for i in range(1, 15)]

                with engine.connect() as conn:
                    _do_existing = pd.read_sql(text(f"""
                        SELECT change_date, new_status FROM {S}.fact_day_off_change
                        WHERE labor_id = :lid AND change_date BETWEEN :start AND :end
                    """), conn, params={"lid": do_labor_id, "start": str(_do_dates[0]), "end": str(_do_dates[-1])})
                _do_override_map = dict(zip(_do_existing["change_date"].astype(str), _do_existing["new_status"]))

                _state_key = f"dayoff_state_{do_labor_id}"
                if _state_key not in st.session_state:
                    st.session_state[_state_key] = {
                        str(d): _do_override_map.get(str(d), _do_default[d.weekday()])
                        for d in _do_dates
                    }
                _do_state = st.session_state[_state_key]

                if st.session_state.get("dayoff_msg"):
                    st.success(st.session_state.pop("dayoff_msg"))

                st.caption("Click a day to toggle between Work and Off, then Save. Only days after today are shown.")

                for week in range(2):
                    cols = st.columns(7)
                    for i in range(7):
                        d      = _do_dates[week * 7 + i]
                        d_key  = str(d)
                        status = _do_state[d_key]
                        with cols[i]:
                            st.caption(d.strftime("%a %m/%d"))
                            label = "🟢 Work" if status == "work" else "🔴 Off"
                            if st.button(label, key=f"dayoff_btn_{do_labor_id}_{d_key}", use_container_width=True):
                                _do_state[d_key] = "off" if status == "work" else "work"
                                st.rerun()
                            if status != _do_default[d.weekday()]:
                                st.caption("modified")

                st.divider()
                c_save, c_discard = st.columns(2)
                if c_save.button("Save Changes", type="primary", key="dayoff_save_btn", use_container_width=True):
                    with engine.begin() as conn:
                        for d in _do_dates:
                            d_key   = str(d)
                            desired = _do_state[d_key]
                            default = _do_default[d.weekday()]
                            if desired == default:
                                conn.execute(text(f"""
                                    DELETE FROM {S}.fact_day_off_change
                                    WHERE labor_id=:lid AND change_date=:cd
                                """), {"lid": do_labor_id, "cd": d_key})
                            else:
                                conn.execute(text(f"""
                                    INSERT INTO {S}.fact_day_off_change (labor_id, change_date, new_status)
                                    VALUES (:lid, :cd, :st)
                                    ON CONFLICT (labor_id, change_date) DO UPDATE SET
                                        new_status = EXCLUDED.new_status, updated_at = NOW()
                                """), {"lid": do_labor_id, "cd": d_key, "st": desired})
                    st.session_state.pop(_state_key, None)
                    st.session_state["dayoff_msg"] = f"Day-off changes saved — **{do_worker_label}**"
                    st.rerun()
                if c_discard.button("Discard Changes", key="dayoff_discard_btn", use_container_width=True):
                    st.session_state.pop(_state_key, None)
                    st.rerun()
